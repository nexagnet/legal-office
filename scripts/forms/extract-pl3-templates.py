#!/usr/bin/env python3
"""Trích văn bản biểu mẫu N-01..N-11 từ Phụ lục III (TT 08/2026/TT-BTP) ra TypeScript.

Chạy MỘT LẦN khi cập nhật mẫu (CI không có tệp nguồn):

    python3 -I scripts/forms/extract-pl3-templates.py <bieu-mau-nghiep-vu-thua-hanh-vien.docx> \
        apps/api/src/forms/templates
    pnpm exec prettier --write apps/api/src/forms/templates

Sinh ra mỗi biểu mẫu một tệp `n-XX.ts`. Văn bản cố định được giữ NGUYÊN; chỗ chấm
`......` ứng với dữ liệu có nguồn được thay bằng `{{khóa=chỗ-chấm-gốc}}`; chỗ không có dữ
liệu giữ nguyên dấu chấm. Số khóa phải khớp số chỗ chấm trong đoạn, nếu không script dừng.
"""
import re
import sys
import unicodedata

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC, OUT = sys.argv[1], sys.argv[2]

RUN_RE = re.compile(r'[.…]{2,}\(\d\)[.…]{2,}|[.…]{2,}|\(năm\)')


def D(prefix):
    """['x.d','x.m','x.y'] cho một ngày tách ba chỗ chấm."""
    return [prefix + '.d', prefix + '.m', prefix + '.y']


SIGN = ['signDay', 'signMonth', 'signYear']
EFF = ['effectiveDate.d', 'effectiveDate.m', 'effectiveDate.y']
SIGN_AT = ['signDate.d', 'signDate.m', 'signDate.y', 'signPlace']
OFFICE_REG = ['office.businessRegNo'] + D('office.businessRegDate')
SIGN_AT = ['signDate.d', 'signDate.m', 'signDate.y', 'signPlace']
ID_CARD = lambda p: [p + '.idNumber']  # noqa: E731
START = ['startHour', 'startMinute', 'startDate.d', 'startDate.m', 'startDate.y', 'place']
END_SAME_DAY = ['endHour', 'endMinute']

# --- phạm vi từng biểu mẫu: (nhãn đầu, nhãn cuối) theo thứ tự thân tài liệu -------------
RANGES = {
    'N-01': ('P7', 'T4'),
    'N-02': ('T6', 'T8'),
    'N-03': ('P128', 'P218'),
    'N-04': ('T12', 'P265'),
    'N-05': ('P266', 'P376'),
    'N-06': ('T18', 'T19'),
    'N-07': ('P408', 'P501'),
    'N-08': ('T23', 'P548'),
    'N-11': ('T29', 'P603'),
}

# --- ánh xạ chỗ chấm -> khóa dữ liệu (None = để trống cho nhân viên điền tay) -------------
MAP = {
    'N-01': {
        'P12': ['contractSeq', 'contractYear'],
        'P19': SIGN_AT,
        'P21': ['office.name'],
        'P22': ['office.address'],
        'P23': ['office.representative'],
        'P24': OFFICE_REG,
        'P25': ['office.bankAccount', 'office.bankName'],
        'P26': ['office.taxCode'],
        'P29': ['requester.courtName', 'requester.procuracyName', 'requester.agencyName'],
        'P30': ['requester.address'],
        'P53': ['payment.days', 'payment.from'],
        'P82': EFF,
        'P97': ['contractSeq', 'contractYear'] + SIGN_AT[:3],
        'T4': SIGN_AT[:3] + ['office.name'],
    },
    'N-02': {
        'T6': ['office.name'],
        'P109': START,
        'P110': ['placeDetail'],
        'P112': ['thv.name'],
        'P113': ['recorder.name', 'recorder.title'],
        'P120': ['recipient.name'],
        'P121': ['recipient.address'],
        'P122': ['recipient.name'],
        'P123': END_SAME_DAY,
    },
    'N-03': {
        'P133': ['contractSeq', 'contractYear'],
        'P140': SIGN_AT,
        'P143': ['office.address'],
        'P144': OFFICE_REG,
        'P145': ['office.bankAccount', 'office.bankName'],
        'P146': ['office.taxCode'],
        'P147': ['office.representative'],
        'P150': ['requester.organization.name'],
        'P151': ['requester.organization.address'],
        'P152': ['requester.organization.representative', 'requester.organization.representativeTitle'],
        'P153': ['requester.organization.bankAccount', 'requester.organization.bankName'],
        'P154': ['requester.organization.taxCode'],
        'P156': ['requester.person.fullName'],
        'P157': ['requester.person.address'],
        'P158': ['requester.person.idNumber'],
        'P159': D('requester.person.idIssueDate') + ['requester.person.idIssuePlace'],
        'P168': ['deadlineDays'] + D('deadlineFrom'),
        'P170': ['fee.amount'],
        'P174': ['fee.total'],
        'P175': ['fee.totalWords'],
        'P177': ['payment.days'] + D('payment.from'),
        'P207': EFF,
    },
    'N-04': {
        'T12': ['office.name', 'vibangSeq', 'vibangYear'],
        'P229': START,
        'P230': ['placeDetail'],
        'P232': ['thv.name'],
        'P233': ['thv.cardNo'] + D('thv.cardIssueDate'),
        'P235': ['requester.name'],
        'P236': ['requester.address'],
        'P237': ['requester.idNumber'],
        'P238': D('requester.idIssueDate') + ['requester.idIssuePlace'],
        'P241': ['participant.name'],
        'P242': ['participant.address'],
        'P243': ['participant.idNumber'],
        'P244': D('participant.idIssueDate') + ['participant.idIssuePlace'],
        'P252': ['eventEnd.hour', 'eventEnd.minute'] + D('eventEnd.date'),
        'P259': ['copies', 'doneHour', 'doneMinute'] + D('doneDate'),
    },
    'N-05': {
        'P271': ['contractSeq', 'contractYear'],
        'P278': ['requesterName', 'requesterAddress'],
        'P280': SIGN_AT,
        'P282': ['office.name'],
        'P283': ['office.address'],
        'P284': ['office.representative'],
        'P285': OFFICE_REG,
        'P286': ['office.bankAccount', 'office.bankName'],
        'P287': ['office.taxCode'],
        'P290': ['requester.organization.name'],
        'P291': ['requester.organization.address'],
        'P292': ['requester.organization.representative', 'requester.organization.representativeTitle'],
        'P293': ['requester.organization.bankAccount', 'requester.organization.bankName'],
        'P294': ['requester.organization.taxCode'],
        'P296': ['requester.person.fullName'],
        'P297': ['requester.person.address'],
        'P298': ['requester.person.idNumber'],
        'P299': D('requester.person.idIssueDate') + ['requester.person.idIssuePlace'],
        'P305': ['subject.person.fullName'],
        'P306': ['subject.person.address'],
        'P307': ['subject.person.idNumber'],
        'P308': D('subject.person.idIssueDate') + ['subject.person.idIssuePlace'],
        'P311': ['subject.organization.name'],
        'P312': ['subject.organization.address'],
        'P313': ['subject.organization.representative', 'subject.organization.representativeTitle'],
        'P319': ['deadlineText'],
        'P327': ['fee.total'],
        'P328': ['fee.totalWords'],
        'P330': ['payment.days', 'payment.from'],
        'P366': EFF,
    },
    'N-06': {
        'T18': ['office.name', 'decisionSeq', 'signPlace'] + D('signDate'),
        'P394': ['judgment.no'] + D('judgment.date') + ['judgment.court'],
        'P395': ['enforcementDecision.no'] + D('enforcementDecision.date') + ['enforcementDecision.office'],
        'P396': ['contract.no'] + D('contract.date') + ['contract.signedWith', 'contract.signedWithAddress'],
        'P400': ['subject.name', 'subject.address'],
        'P404': ['assignedThv.name'],
    },
    'N-07': {
        'P411': ['contractSeq', 'contractYear'],
        'P418': ['requesterName', 'requesterAddress'],
        'P420': SIGN_AT,
        'P422': ['office.name'],
        'P423': ['office.address'],
        'P424': OFFICE_REG,
        'P425': ['office.bankAccount', 'office.bankName'],
        'P426': ['office.taxCode'],
        'P427': ['office.representative'],
        'P430': ['requester.organization.name'],
        'P431': ['requester.organization.address'],
        'P432': ['requester.organization.representative', 'requester.organization.representativeTitle'],
        'P433': ['requester.organization.bankAccount', 'requester.organization.bankName'],
        'P434': ['requester.organization.taxCode'],
        'P436': ['requester.person.fullName'],
        'P437': ['requester.person.address'],
        'P438': ['requester.person.idNumber'],
        'P439': D('requester.person.idIssueDate') + ['requester.person.idIssuePlace'],
        'P447': ['judgment.no'] + D('judgment.date') + ['judgment.court'],
        'P448': ['deadlineMonths'],
        'P450': ['fee.verification'],
        'P452': ['fee.percent', 'fee.minimum', 'office.name'],
        'P453': ['fee.advance'],
        'P462': ['payment.days', 'payment.from'],
        'P493': EFF,
    },
    'N-08': {
        'T23': ['office.name', 'signPlace'] + D('signDate'),
        'P507': START,
        'P511': ['giver.office'],
        'P510': ['giver.name'],
        'P513': ['receiver.name', 'receiver.title'],
        'P514': ['receiver.organization'],
        'P516': ['justiceDept.name', 'justiceDept.title'],
        'P517': ['justiceDept.province'],
        'P520': ['fileCount'],
        'P521': ['fileCondition'],
        'P523': ['documents'],
        'P524': ['accountingVouchers'],
        'P525': ['otherDocuments'],
        'P527': [None, 'storageDevices'],
        'P528': ['systemData'],
        'P535': ['receiver.name'],
        'P536': ['endHour', 'endMinute', 'copies.giver', 'copies.receiver', 'copies.justiceDept', 'copies.file'],
    },
    'N-11': {
        'T29': ['office.name', 'signPlace'] + D('signDate'),
        'P579': [None, None] + D('vibang.date'),
        'P580': ['office.name'],
        'P582': ['vibang.number'],
        'P583': ['vibang.date'],
        'P584': ['thv.name', 'thv.cardNo'],
        'P585': ['requester.name'],
        'P590': ['copies', None, None] + D('vibang.date'),
        'P593': D('handedDate'),
    },
}

# Thay thế văn bản chính xác SAU khi gán khóa (chỗ không có dấu chấm trong bản gốc).
SPECIAL = {
    'N-01': {
        'P31': [('Đại diện: ông (bà) , chức vụ:',
                 'Đại diện: ông (bà) {{requester.representative=......}}, chức vụ: {{requester.representativeTitle=......}}')],
        'P32': [('Số tài khoản: tại Kho bạc Nhà nước',
                 'Số tài khoản: {{requester.bankAccount=......}} tại Kho bạc Nhà nước')],
        'P33': [('Mã số thuế:', 'Mã số thuế: {{requester.taxCode=......}}')],
    },
    'N-03': {
        'P142': [('**1. Văn phòng thi hành án dân sự**:',
                  '**1. Văn phòng thi hành án dân sự**: {{office.name=}}')],
    },
    'N-08': {
        'P510': [('Thừa hành viên/Trưởng văn phòng',
                  '{{giver.title=Thừa hành viên/Trưởng văn phòng}}')],
    },
    'N-11': {
        'T31': [('THỪA HÀNH VIÊN/TRƯỞNG VĂN PHÒNG', '{{signerTitle=THỪA HÀNH VIÊN/TRƯỞNG VĂN PHÒNG}}')],
    },
}

# Nhóm đoạn chỉ gồm dòng chấm thành khối `lines` (nhiều dòng văn bản tự do).
LINES = {
    'N-03': {'P165': ('subjectMatter', 2)},
    'N-04': {'P247': ('eventTitle', 2), 'P250': ('eventContent', 2)},
    'N-05': {'P316': ('verificationContent', 3)},
    'N-06': {'P402': ('verificationMethod', 2)},
    'N-07': {'P445': ('enforcementSubject', 2)},
    'N-08': {'P530': ('otherFiles', 2), 'P533': ('opinions', 2)},
    'N-11': {'P588': ('reason', 2)},
}
NUMBERED = {
    'N-02': {'P118': ('documents', 2)},
    'N-04': {'P255': ('attachments', 2)},
}
# Bảng động thay bằng khối riêng.
DYNAMIC_TABLES = {'T3': 'itemsTable', 'T30': 'correctionsTable'}
# Bảng mã/ký hiệu mẫu và các bảng thuộc mẫu khác: bỏ khỏi thân.
DROP = {'T1', 'T5', 'T9', 'T11', 'T15', 'T17', 'T20', 'T22', 'T28'}

JOIN_FIXES = [('tổ\nchức', 'tổ chức'), ('cá\nnhân', 'cá nhân'), ('LẬP VI\nBẰNG', 'LẬP VI BẰNG')]


def nfc(s):
    return unicodedata.normalize('NFC', s)


def esc(s):
    return s.replace('\\', '\\\\').replace("'", "\\'")


def run_markup(p):
    """Chuỗi có đánh dấu **đậm** và %%nghiêng%% từ các run của đoạn."""
    out = []
    state = (False, False)
    for r in p.runs:
        t = r.text
        if t == '':
            continue
        b, i = bool(r.bold), bool(r.italic)
        if t.strip() == '':
            out.append((t, state[0], state[1], True))
            continue
        out.append((t, b, i, False))
    # gộp run liên tiếp cùng định dạng, đẩy khoảng trắng ra ngoài dấu
    s = ''
    cur = None
    buf = ''

    def flush():
        nonlocal s, buf, cur
        if buf == '':
            return
        core = buf.strip(' ')
        lead = buf[: len(buf) - len(buf.lstrip(' '))]
        trail = buf[len(buf.rstrip(' ')):]
        if core == '' or cur == (False, False):
            s += buf
        else:
            def wrap(line):
                if line == '':
                    return line
                if cur[1]:
                    line = '%%' + line + '%%'
                if cur[0]:
                    line = '**' + line + '**'
                return line

            s += lead + '\n'.join(wrap(x) for x in core.split('\n')) + trail
        buf = ''

    for t, b, i, ws in out:
        fmt = cur if ws else (b, i)
        if fmt != cur:
            flush()
            cur = fmt
        buf += t
    flush()
    return s


def assign(text, keys, label):
    runs = list(RUN_RE.finditer(text))
    if len(runs) != len(keys):
        raise SystemExit(f'{label}: {len(runs)} chỗ chấm nhưng {len(keys)} khóa: {text!r}')
    out, last = '', 0
    for m, k in zip(runs, keys):
        out += text[last:m.start()]
        if k is None:
            out += m.group(0)
        else:
            # bản gốc dính chữ với chấm ("hồi........"): tách bằng dấu cách để giá trị điền
            # vào không dính chữ liền kề ("hồi09").
            if out and (out[-1].isalnum() or out[-1] in ':)'):
                out += ' '
            out += '{{%s=%s}}' % (k, m.group(0))
            nxt = text[m.end():m.end() + 1]
            if nxt and nxt.isalnum():
                out += ' '
            elif re.match(r'^\s+[,;:.]', text[m.end():]):
                # bản gốc có khoảng trắng thừa trước dấu câu ("năm ...... , tại")
                last = m.end() + len(text[m.end():]) - len(text[m.end():].lstrip())
                continue
        last = m.end()
    return out + text[last:]


def clean(s):
    s = nfc(s)
    s = re.sub(r'[ \t ]+', ' ', s)
    return s.strip(' ')


def para_text(p):
    return run_markup(p)


def align_of(p):
    a = p.alignment
    if a is None:
        return 'left'
    return {1: 'center', 3: 'justify', 2: 'right', 0: 'left'}.get(int(a), 'left')


d = docx.Document(SRC)
body = []
pi = ti = 0
for el in d.element.body.iterchildren():
    tag = el.tag.split('}')[1]
    if tag == 'p':
        body.append(('P%d' % pi, Paragraph(el, d)))
        pi += 1
    elif tag == 'tbl':
        body.append(('T%d' % ti, Table(el, d)))
        ti += 1
labels = [b[0] for b in body]


def fmt_lines(text):
    return [clean(x) for x in text.split('\n')]


def cell_lines(cell, keys_state, label, mapping):
    raw = '\n'.join(run_markup(p) for p in cell.paragraphs)
    for a, b in JOIN_FIXES:
        raw = raw.replace(a, b)
    return raw


def emit_table(label, t, form, keys):
    rows = t.rows
    seen = set()
    cells = []
    for c in rows[0].cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        cells.append(c)
    # nối văn bản toàn bảng để gán khóa theo thứ tự
    texts = []
    for c in cells:
        raw = '\n'.join(run_markup(p) for p in c.paragraphs)
        for a, b in JOIN_FIXES:
            raw = raw.replace(a, b)
        texts.append(nfc(raw))
    if label in MAP.get(form, {}):
        joined = '\u0001'.join(texts)
        joined = assign(joined, MAP[form][label], form + ' ' + label)
        texts = joined.split('\u0001')
    for old, new in SPECIAL.get(form, {}).get(label, []):
        hit = False
        for i, tx in enumerate(texts):
            if old in tx:
                texts[i] = tx.replace(old, new)
                hit = True
        if not hit:
            raise SystemExit(f'{form} {label}: không thấy {old!r}')
    parts = []
    for c, tx in zip(cells, texts):
        lines = [re.sub(r'[ \t ]+', ' ', x).strip() for x in tx.split('\n')]
        while lines and lines[-1] == '':
            lines.pop()
        al = align_of(c.paragraphs[0]) if c.paragraphs else 'center'
        parts.append("cell([%s], '%s')" % (', '.join("'%s'" % esc(x) for x in lines), al))
    return 'cols(%s)' % ', '.join(parts)


def emit_form(form):
    first, last = RANGES[form]
    last = last
    i0 = labels.index(first)
    i1 = labels.index(last)
    els = body[i0:i1 + 1]
    # N-01: phụ lục Bảng kê nằm sau Ghi chú: phạm vi đã gồm đến T4.
    out = []
    mapping = MAP.get(form, {})
    special = SPECIAL.get(form, {})
    lines_blocks = LINES.get(form, {})
    numbered = NUMBERED.get(form, {})
    skip_until = None
    k = 0
    while k < len(els):
        label, el = els[k]
        k += 1
        if label in DROP:
            continue
        if label.startswith('T'):
            if label in DYNAMIC_TABLES:
                out.append(DYNAMIC_TABLES[label] + '()')
                continue
            t = el
            if len(t.rows) == 1:
                out.append(emit_table(label, t, form, None))
                continue
            raise SystemExit(f'{form} {label}: bảng nhiều dòng chưa hỗ trợ')
        # đoạn
        p = el
        if label in lines_blocks or label in numbered:
            key, n = (lines_blocks.get(label) or numbered.get(label))
            blanks = []
            for j in range(n):
                _, pj = els[k - 1 + j]
                tx = clean(pj.text)
                blanks.append(tx)
            k += n - 1
            if label in lines_blocks:
                for b in blanks:
                    assert re.fullmatch(r'[.…]+', b), (label, b)
                out.append("lines('%s', [%s])" % (key, ', '.join("'%s'" % esc(b) for b in blanks)))
            else:
                out.append("numbered('%s', [%s])" % (key, ', '.join("'%s'" % esc(b) for b in blanks)))
            continue
        tx = nfc(para_text(p))
        if tx.strip() == '':
            # đoạn trống: giữ một dòng cách duy nhất trong thân
            if out and out[-1] == 'blank()':
                continue
            out.append('blank()')
            continue
        tx = re.sub(r'[ \t ]+', ' ', tx).strip()
        if label in mapping:
            tx = assign(tx, mapping[label], form + ' ' + label)
        for old, new in special.get(label, []):
            if old not in tx:
                raise SystemExit(f'{form} {label}: không thấy {old!r} trong {tx!r}')
            tx = tx.replace(old, new)
        # Điều ngắn: in đậm cả dòng
        plain = tx.replace('**', '')
        if re.match(r'^Điều \d+\.', plain) and len(plain) < 80 and '{{' not in plain:
            tx = '**' + plain + '**'
        al = align_of(p)
        fn = {'center': 'c', 'justify': 'j', 'left': 'l', 'right': 'r'}[al]
        opt = ''
        parts = tx.split('\n')
        for q in parts:
            out.append("%s('%s'%s)" % (fn, esc(q), opt))
    while out and out[-1] == 'blank()':
        out.pop()
    while out and out[0] == 'blank()':
        out.pop(0)
    return out


NAMES = {
    'N-01': ('n-01', 'N01_TEMPLATE', 'Hợp đồng dịch vụ tống đạt'),
    'N-02': ('n-02', 'N02_TEMPLATE', 'Biên bản thực hiện việc tống đạt'),
    'N-03': ('n-03', 'N03_TEMPLATE', 'Hợp đồng dịch vụ lập vi bằng'),
    'N-04': ('n-04', 'N04_TEMPLATE', 'Vi bằng'),
    'N-05': ('n-05', 'N05_TEMPLATE', 'Hợp đồng dịch vụ xác minh điều kiện thi hành án'),
    'N-06': ('n-06', 'N06_TEMPLATE', 'Quyết định xác minh điều kiện thi hành án'),
    'N-07': ('n-07', 'N07_TEMPLATE', 'Hợp đồng dịch vụ thi hành án'),
    'N-08': ('n-08', 'N08_TEMPLATE', 'Biên bản giao nhận hồ sơ, tài liệu'),
    'N-11': ('n-11', 'N11_TEMPLATE', 'Văn bản sửa lỗi kỹ thuật vi bằng'),
}

for form in RANGES:
    fname, const, title = NAMES[form]
    blocks = emit_form(form)
    used = sorted(
        {m for b in blocks for m in re.findall(r'\{\{([\w.]+)=', b)}
    )
    helpers = {b.split('(')[0] for b in blocks}
    if 'cols' in helpers:
        helpers.add('cell')
    helpers = sorted(helpers)
    src = (
        '// Sinh tự động từ Phụ lục III, Thông tư 08/2026/TT-BTP bởi\n'
        '// scripts/forms/extract-pl3-templates.py. Sửa văn bản mẫu: chạy lại script, không sửa tay.\n'
        "import { %s, type FormTemplate } from '../form-template';\n\n"
        'export const %s: FormTemplate = {\n'
        "  code: '%s',\n"
        "  ref: 'TP-THV-%s',\n"
        "  title: '%s',\n"
        '  blocks: [\n%s\n  ],\n'
        '};\n'
    ) % (
        ', '.join(helpers),
        const,
        form,
        form,
        esc(title),
        ',\n'.join('    ' + b for b in blocks),
    )
    with open(f'{OUT}/{fname}.ts', 'w', encoding='utf8') as f:
        f.write(src)
    print(form, len(blocks), 'khối;', len(used), 'khóa')
